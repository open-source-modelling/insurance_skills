"""Generates every example workbook in ``examples/``.

Binary .xlsx files cannot be reviewed in a pull request — but this script can,
so a reader can see exactly what makes each bad file bad. Both the script and
its output are committed, so nobody has to run anything to open an example.

Each bad fixture carries **exactly one** defect, and declares which rule it is
meant to trip. ``tests/test_fixtures.py`` asserts that the runner reports that
rule and only that rule, which catches false negatives and false positives in
the same assertion.

Usage::

    python scripts/make_fixtures.py
"""

from __future__ import annotations

import datetime as dt
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent.parent  # the skill root
GOOD_DIR = ROOT / "examples" / "good"
BAD_DIR = ROOT / "examples" / "bad"

DATE_FORMAT = "dd/mm/yyyy"
DOWNLOAD_TIME = dt.datetime(2025, 12, 31)
SHEETS = ("FUND_1", "FUND_2", "FUND_3")

HEADERS = [
    "DOWNLOAD_TIME", "START_DATE", "FUND_NAME", "POLICY_ID", "STATUS",
    "GENDER", "BIRTH_DATE", "NUM_UNITS", "PRICE_UNIT", "ANNUAL_PREMIUM",
]

# Four rows chosen to exercise the edges the contract explicitly allows:
#   - ANNUAL_PREMIUM = 0 on an ACTIVE policy, and positive on a PASSIVE one
#   - NUM_UNITS = 0
#   - the same POLICY_ID reused on every sheet (allowed; unique per sheet only)
BASE_ROWS = [
    dict(START_DATE=dt.datetime(2020, 12, 31), POLICY_ID=1234556, STATUS="ACTIVE",
         GENDER="M", BIRTH_DATE=dt.datetime(1978, 1, 12),
         NUM_UNITS=1234, PRICE_UNIT=12.5, ANNUAL_PREMIUM=4532),
    dict(START_DATE=dt.datetime(2021, 1, 12), POLICY_ID=1234552, STATUS="PASSIVE",
         GENDER="F", BIRTH_DATE=dt.datetime(2000, 2, 15),
         NUM_UNITS=123, PRICE_UNIT=12.5, ANNUAL_PREMIUM=980),
    dict(START_DATE=dt.datetime(2023, 6, 1), POLICY_ID=1234570, STATUS="ACTIVE",
         GENDER="F", BIRTH_DATE=dt.datetime(1995, 11, 3),
         NUM_UNITS=0, PRICE_UNIT=9.75, ANNUAL_PREMIUM=0),
    dict(START_DATE=dt.datetime(2025, 3, 17), POLICY_ID=1234588, STATUS="PASSIVE",
         GENDER="M", BIRTH_DATE=dt.datetime(1960, 8, 22),
         NUM_UNITS=57.5, PRICE_UNIT=11.2, ANNUAL_PREMIUM=1200),
]

DATE_COLUMNS = {"DOWNLOAD_TIME", "START_DATE", "BIRTH_DATE"}


def build_clean() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEETS:
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)
        for row in BASE_ROWS:
            values = {"DOWNLOAD_TIME": DOWNLOAD_TIME, "FUND_NAME": sheet_name, **row}
            ws.append([values[h] for h in HEADERS])
        _apply_date_formats(ws)
    return wb


def _apply_date_formats(ws: Worksheet) -> None:
    for index, header in enumerate(HEADERS, start=1):
        if header in DATE_COLUMNS:
            for row in range(2, ws.max_row + 1):
                ws.cell(row, index).number_format = DATE_FORMAT


def col(header: str) -> int:
    return HEADERS.index(header) + 1


# --------------------------------------------------------------------------
# The mutations. One per rule; each breaks exactly one thing.
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Fixture:
    filename: str
    rule: str
    note: str
    mutate: Callable[[Workbook], None]
    needs_recalc: bool = False


def _drop_column(wb: Workbook, header: str) -> None:
    for ws in wb.worksheets:
        ws.delete_cols(col(header))


def _add_column(wb: Workbook, header: str) -> None:
    for ws in wb.worksheets:
        ws.cell(1, len(HEADERS) + 1, header)
        for row in range(2, ws.max_row + 1):
            ws.cell(row, len(HEADERS) + 1, "x")


def _set(wb: Workbook, sheet: str, header: str, row: int, value, number_format=None) -> None:
    cell = wb[sheet].cell(row, col(header), value)
    if number_format:
        cell.number_format = number_format


def _blank(wb: Workbook, sheet: str, header: str, row: int) -> None:
    """Empty a cell. ``ws.cell(..., value=None)`` is a no-op in openpyxl."""
    wb[sheet].cell(row, col(header)).value = None


FIXTURES = [
    Fixture("w1_extra_sheet.xlsx", "W1", "a fourth sheet nobody expects",
            lambda wb: wb.create_sheet("FUND_4").append(HEADERS)),
    Fixture("w2_header_order.xlsx", "W2", "FUND_2 has GENDER and STATUS swapped",
            lambda wb: _swap_headers(wb, "FUND_2", "STATUS", "GENDER")),
    Fixture("w3_blank_row.xlsx", "W3", "a blank row in the middle of FUND_1",
            lambda wb: wb["FUND_1"].insert_rows(3)),
    Fixture("w4_merged_cells.xlsx", "W4", "a merged region left beside the data on FUND_3",
            lambda wb: wb["FUND_3"].merge_cells("L1:M1")),
    Fixture("c1_missing_column.xlsx", "C1", "GENDER dropped from every sheet",
            lambda wb: _drop_column(wb, "GENDER")),
    Fixture("c2_extra_column.xlsx", "C2", "an undocumented NOTES column",
            lambda wb: _add_column(wb, "NOTES")),
    Fixture("f1_formula_dates.xlsx", "F1", "BIRTH_DATE written as =DATE(...) formulas",
            lambda wb: _formula_dates(wb), needs_recalc=True),
    Fixture("f2_text_date.xlsx", "F2", "one START_DATE stored as the text '31/12/2020'",
            lambda wb: _set(wb, "FUND_1", "START_DATE", 2, "31/12/2020", DATE_FORMAT)),
    Fixture("f3_mixed_date_format.xlsx", "F3", "FUND_2 uses mm/dd/yyyy for START_DATE",
            lambda wb: _reformat(wb, "FUND_2", "START_DATE", "mm/dd/yyyy")),
    Fixture("d1_two_download_times.xlsx", "D1", "FUND_3 was extracted a day later",
            lambda wb: _set(wb, "FUND_3", "DOWNLOAD_TIME", 2,
                            dt.datetime(2026, 1, 1), DATE_FORMAT)),
    Fixture("d3_fund_name_mismatch.xlsx", "D3", "a FUND_1 row labelled FUND_2",
            lambda wb: _set(wb, "FUND_1", "FUND_NAME", 3, "FUND_2")),
    Fixture("d4_policy_id_as_text.xlsx", "D4", "a POLICY_ID stored as text",
            lambda wb: _set(wb, "FUND_2", "POLICY_ID", 2, "1234556")),
    Fixture("d5_duplicate_policy_id.xlsx", "D5", "the same POLICY_ID twice on FUND_1",
            lambda wb: _set(wb, "FUND_1", "POLICY_ID", 3, 1234556)),
    Fixture("d6_bad_status.xlsx", "D6", "STATUS of CLOSED",
            lambda wb: _set(wb, "FUND_1", "STATUS", 2, "CLOSED")),
    Fixture("d7_blank_gender.xlsx", "D7", "a blank GENDER",
            lambda wb: _blank(wb, "FUND_3", "GENDER", 4)),
    Fixture("d8_implausible_birth_date.xlsx", "D8", "a BIRTH_DATE in 1899",
            lambda wb: _birth_and_start(wb, dt.datetime(1899, 6, 1), dt.datetime(1980, 1, 1))),
    Fixture("d9_negative_units.xlsx", "D9", "a negative NUM_UNITS",
            lambda wb: _set(wb, "FUND_2", "NUM_UNITS", 3, -5)),
    Fixture("d10_zero_price.xlsx", "D10", "a PRICE_UNIT of exactly 0",
            lambda wb: _set(wb, "FUND_1", "PRICE_UNIT", 4, 0)),
    Fixture("d11_negative_premium.xlsx", "D11", "a negative ANNUAL_PREMIUM",
            lambda wb: _set(wb, "FUND_3", "ANNUAL_PREMIUM", 2, -100)),
    Fixture("x1_born_after_start.xlsx", "X1", "BIRTH_DATE later than START_DATE",
            lambda wb: _set(wb, "FUND_1", "BIRTH_DATE", 2,
                            dt.datetime(2021, 5, 1), DATE_FORMAT)),
    Fixture("x2_start_after_download.xlsx", "X2", "a policy starting after the extract",
            lambda wb: _set(wb, "FUND_2", "START_DATE", 4,
                            dt.datetime(2026, 2, 1), DATE_FORMAT)),
    Fixture("x3_age_over_limit.xlsx", "X3", "a policyholder aged 123 at START_DATE",
            lambda wb: _birth_and_start(wb, dt.datetime(1901, 1, 1), dt.datetime(2024, 1, 1))),
]


def _swap_headers(wb: Workbook, sheet: str, a: str, b: str) -> None:
    ws = wb[sheet]
    ia, ib = col(a), col(b)
    for row in range(1, ws.max_row + 1):
        va, vb = ws.cell(row, ia).value, ws.cell(row, ib).value
        ws.cell(row, ia, vb)
        ws.cell(row, ib, va)


def _reformat(wb: Workbook, sheet: str, header: str, number_format: str) -> None:
    ws = wb[sheet]
    for row in range(2, ws.max_row + 1):
        ws.cell(row, col(header)).number_format = number_format


def _formula_dates(wb: Workbook) -> None:
    for ws in wb.worksheets:
        for row in range(2, ws.max_row + 1):
            date = BASE_ROWS[row - 2]["BIRTH_DATE"]
            cell = ws.cell(row, col("BIRTH_DATE"),
                           f"=DATE({date.year},{date.month},{date.day})")
            cell.number_format = DATE_FORMAT


def _birth_and_start(wb: Workbook, birth: dt.datetime, start: dt.datetime) -> None:
    """Set both dates on one row, so only the intended rule trips."""
    _set(wb, "FUND_1", "BIRTH_DATE", 2, birth, DATE_FORMAT)
    _set(wb, "FUND_1", "START_DATE", 2, start, DATE_FORMAT)


# --------------------------------------------------------------------------


def _recalc(path: Path) -> None:
    """Bake cached values into formula cells.

    openpyxl writes formulas with no cached result, so without this a formula
    fixture would also fail the "is it a real date" rule — two defects instead
    of the one it is meant to demonstrate.
    """
    script = Path("/mnt/skills/public/xlsx/scripts/recalc.py")
    if not script.exists():
        print(f"  ! recalc script unavailable; {path.name} will have empty formula cells")
        return
    subprocess.run([sys.executable, str(script), str(path)], check=False,
                   capture_output=True)


def main() -> None:
    GOOD_DIR.mkdir(parents=True, exist_ok=True)
    BAD_DIR.mkdir(parents=True, exist_ok=True)

    clean_path = GOOD_DIR / "UL_FUND_DATA_Q4_2025.xlsx"
    build_clean().save(clean_path)
    print(f"wrote {clean_path.relative_to(ROOT)}")

    for fixture in FIXTURES:
        wb = build_clean()
        fixture.mutate(wb)
        path = BAD_DIR / fixture.filename
        wb.save(path)
        if fixture.needs_recalc:
            _recalc(path)
        print(f"wrote {path.relative_to(ROOT)}  [{fixture.rule}] {fixture.note}")


if __name__ == "__main__":
    main()
